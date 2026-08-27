# -*- coding: utf-8 -*-
# MAME XML → XLSX 转换脚本
# 将 MAME 的 XML 数据库（如 mame0288.xml）转换为 Excel 表格。
# 支持两种输出方式：
#   1. 全部数据合并为一个 XLSX 文件
#   2. 按类别拆分为四个 XLSX 文件（Arcade ROM / Arcade CHD / Software ROM / Software CHD）
# 用法：python MameXml2Xslx.py [xml文件] [输出文件(可选)]
# 不带参数时交互式询问 XML 位置（默认 d:\ProApps\mame\current\mame0288.xml）。

import os
import signal
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# ==================== 全局配置 ====================
DEFAULT_MAME_XML = Path(r"d:\ProApps\mame\current\mame0288.xml")
MIN_COL_WIDTH = 10              # 列宽下限
MAX_COL_WIDTH = 60              # 列宽上限（防 SHA1/sourcefile 等长文本撑爆）
COL_WIDTH_SAMPLE_ROWS = 200     # 列宽计算采样行数（避免全表遍历，大表性能优化）

# --- 消息常量：提示 ---
MSG_PROMPT_XML = "请输入 XML 文件位置"
MSG_PROMPT_MODE = "请选择 [1/2]: "
MSG_TITLE = "MAME XML → XLSX"
MSG_READING_XML = "读取 MAME XML"
MSG_PARSING = "正在解析，请稍候..."
MSG_PARSING_ARCADE = "解析 Arcade Machines..."
MSG_PARSING_SOFTWARE = "解析 Software Lists..."
MSG_ADJUST_WIDTH = "调整列宽..."
MSG_SEPARATOR = "=" * 60
MSG_LINE = "-" * 60
MSG_XML_PATH = "文件: {}"

# --- 消息常量：状态 ---
MSG_PARSE_DONE = "解析完成"
MSG_GENERATING = "生成: {}"
MSG_DONE = "完成: {}"
MSG_ALL_DONE = "全部完成"
MSG_SKIP_EMPTY = "跳过空数据集: {}"
MSG_STAT_ARCADE_ROM = "Arcade ROM       : {:,}"
MSG_STAT_ARCADE_CHD = "Arcade CHD       : {:,}"
MSG_STAT_SOFTWARE_ROM = "Software ROM     : {:,}"
MSG_STAT_SOFTWARE_CHD = "Software CHD     : {:,}"
MSG_STAT_TOTAL = "Total            : {:,}"
MSG_MODE_MENU = """请选择输出方式：

1. 一个 XLSX 文件
   └─ Arcade ROM + Arcade CHD + Software ROM + Software CHD

2. 四个 XLSX 文件
   ├─ Arcade ROM
   ├─ Arcade CHD
   ├─ Software ROM
   └─ Software CHD
"""

# --- 消息常量：错误 / 退出 ---
MSG_INVALID_CHOICE = "输入错误，请输入 1 或 2。"
MSG_FILE_NOT_FOUND = "错误：找不到文件：{}"
MSG_NO_DATA = "错误：未解析到任何数据。"
MSG_INTERRUPTED = "\n\n用户中断程序，已退出。"
MSG_ERROR = "\n程序运行出错: {}"
MSG_EXIT = "\n按回车键退出..."
MSG_USAGE = """用法:
    python {script_name} mame0288.xml
    或
    python {script_name} mame0288.xml output.xlsx
不带参数时交互式询问 XML 文件位置。"""


# ==================== 中断处理 ====================
_quit_requested = False


def _on_quit_signal(signum, frame):
    global _quit_requested
    _quit_requested = True
    raise KeyboardInterrupt()


def _init_quit_handler():
    if hasattr(signal, "SIGQUIT"):
        signal.signal(signal.SIGQUIT, _on_quit_signal)


def _check_quit() -> bool:
    global _quit_requested
    if sys.platform == "win32":
        try:
            import msvcrt
            while msvcrt.kbhit():
                if msvcrt.getch() == b"\x11":
                    _quit_requested = True
        except Exception:
            pass
    return _quit_requested


# ==================== Excel 列定义 ====================
FIELDS = [
    "category",
    "machine",
    "game_name",
    "software_list",
    "software",
    "software_description",
    "part",
    "interface",
    "type",
    "filename",
    "size",
    "crc",
    "sha1",
    "md5",
    "merge",
    "region",
    "offset",
    "status",
    "optional",
    "machine_cloneof",
    "machine_romof",
    "sourcefile",
]


# ==================== XML 读取辅助 ====================
def get_attr(element, name, default=""):
    """读取元素属性；元素不存在时返回默认值。"""
    if element is None:
        return default
    return element.attrib.get(name, default)


def get_text(element) -> str:
    """读取元素文本（如 <description>xxx</description>）。"""
    if element is None or element.text is None:
        return ""
    return element.text.strip()


def make_row(**kwargs):
    """按 FIELDS 顺序创建一条记录，未提供的字段填空串。"""
    return {
        field: kwargs.get(field, "")
        for field in FIELDS
    }


# ==================== 解析 MAME XML ====================
def parse_mame_xml(xml_file):
    """解析 MAME XML，返回记录列表。统计信息在函数内打印。"""

    print()
    print(MSG_SEPARATOR)
    print(MSG_READING_XML)
    print(MSG_SEPARATOR)
    print(MSG_XML_PATH.format(xml_file))
    print()
    print(MSG_PARSING)

    tree = ET.parse(xml_file)
    root = tree.getroot()

    rows = []
    arcade_rom_count = 0
    arcade_chd_count = 0
    software_rom_count = 0
    software_chd_count = 0

    # ========================================================
    # Arcade / MAME Machines
    # ========================================================

    print(MSG_PARSING_ARCADE)

    for machine in root.findall("machine"):

        if _check_quit():
            raise KeyboardInterrupt()

        machine_name = get_attr(machine, "name")
        description = get_text(machine.find("description"))

        cloneof = get_attr(machine, "cloneof")
        romof = get_attr(machine, "romof")
        sourcefile = get_attr(machine, "sourcefile")

        # ----------------------------------------------------
        # Arcade ROM
        # ----------------------------------------------------

        for rom in machine.findall("rom"):

            row = make_row(
                category="arcade",
                machine=machine_name,
                game_name=description,
                type="rom",
                filename=get_attr(rom, "name"),
                size=get_attr(rom, "size"),
                crc=get_attr(rom, "crc"),
                sha1=get_attr(rom, "sha1"),
                md5=get_attr(rom, "md5"),
                merge=get_attr(rom, "merge"),
                region=get_attr(rom, "region"),
                offset=get_attr(rom, "offset"),
                status=get_attr(rom, "status"),
                optional=get_attr(rom, "optional"),
                machine_cloneof=cloneof,
                machine_romof=romof,
                sourcefile=sourcefile,
            )

            rows.append(row)
            arcade_rom_count += 1

        # ----------------------------------------------------
        # Arcade CHD
        # ----------------------------------------------------

        for disk in machine.findall("disk"):

            row = make_row(
                category="arcade",
                machine=machine_name,
                game_name=description,
                type="chd",
                filename=get_attr(disk, "name"),
                sha1=get_attr(disk, "sha1"),
                merge=get_attr(disk, "merge"),
                status=get_attr(disk, "status"),
                optional=get_attr(disk, "optional"),
                machine_cloneof=cloneof,
                machine_romof=romof,
                sourcefile=sourcefile,
            )

            rows.append(row)
            arcade_chd_count += 1

    # ========================================================
    # Software Lists
    # ========================================================

    print(MSG_PARSING_SOFTWARE)

    for softwarelist in root.findall("softwarelist"):

        if _check_quit():
            raise KeyboardInterrupt()

        list_name = get_attr(softwarelist, "name")
        list_description = get_attr(softwarelist, "description")

        for software in softwarelist.findall("software"):

            software_name = get_attr(software, "name")
            software_description = get_text(software.find("description"))

            # ------------------------------------------------
            # Software parts
            # ------------------------------------------------

            for part in software.findall("part"):

                part_name = get_attr(part, "name")
                interface = get_attr(part, "interface")

                # ============================================
                # Software ROM
                # ============================================

                for dataarea in part.findall("dataarea"):

                    for rom in dataarea.findall("rom"):

                        row = make_row(
                            category="software",
                            game_name=software_description,
                            software_list=list_name,
                            software=software_name,
                            software_description=software_description,
                            part=part_name,
                            interface=interface,
                            type="rom",
                            filename=get_attr(rom, "name"),
                            size=get_attr(rom, "size"),
                            crc=get_attr(rom, "crc"),
                            sha1=get_attr(rom, "sha1"),
                            md5=get_attr(rom, "md5"),
                            merge=get_attr(rom, "merge"),
                            region=get_attr(rom, "region"),
                            offset=get_attr(rom, "offset"),
                            status=get_attr(rom, "status"),
                            optional=get_attr(rom, "optional"),
                        )

                        rows.append(row)
                        software_rom_count += 1

                # ============================================
                # Software CHD
                # ============================================

                for diskarea in part.findall("diskarea"):

                    for disk in diskarea.findall("disk"):

                        row = make_row(
                            category="software",
                            game_name=software_description,
                            software_list=list_name,
                            software=software_name,
                            software_description=software_description,
                            part=part_name,
                            interface=interface,
                            type="chd",
                            filename=get_attr(disk, "name"),
                            sha1=get_attr(disk, "sha1"),
                            merge=get_attr(disk, "merge"),
                            status=get_attr(disk, "status"),
                            optional=get_attr(disk, "optional"),
                        )

                        rows.append(row)
                        software_chd_count += 1

    # ========================================================
    # 统计
    # ========================================================

    print()
    print(MSG_SEPARATOR)
    print(MSG_PARSE_DONE)
    print(MSG_SEPARATOR)
    print(MSG_STAT_ARCADE_ROM.format(arcade_rom_count))
    print(MSG_STAT_ARCADE_CHD.format(arcade_chd_count))
    print(MSG_STAT_SOFTWARE_ROM.format(software_rom_count))
    print(MSG_STAT_SOFTWARE_CHD.format(software_chd_count))
    print(MSG_LINE)
    print(MSG_STAT_TOTAL.format(len(rows)))
    print()

    return rows


# ==================== 创建 XLSX ====================
def create_xlsx(rows, output_file, sheet_name="MAME Database"):
    """将记录列表写入 XLSX。先写临时文件再原子替换，防中断留下损坏文件。"""

    print(MSG_GENERATING.format(output_file))

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    # --------------------------------------------------------
    # Header
    # --------------------------------------------------------

    ws.append(FIELDS)

    for cell in ws[1]:
        cell.font = Font(bold=True)

    # --------------------------------------------------------
    # 数据
    # --------------------------------------------------------

    for row in rows:
        ws.append([row[field] for field in FIELDS])

    # --------------------------------------------------------
    # 冻结第一行
    # --------------------------------------------------------

    ws.freeze_panes = "A2"

    # --------------------------------------------------------
    # 自动筛选 / Excel Table
    # --------------------------------------------------------

    if rows:

        last_row = len(rows) + 1
        last_column = get_column_letter(len(FIELDS))

        table_ref = f"A1:{last_column}{last_row}"

        table = Table(displayName="MAME_Database", ref=table_ref)
        style = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )

        table.tableStyleInfo = style
        ws.add_table(table)

    # --------------------------------------------------------
    # 自动列宽（只采样前 N 行，几十万行的大表避免全表遍历）
    # --------------------------------------------------------

    print(MSG_ADJUST_WIDTH)

    for column_cells in ws.columns:

        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)

        for cell in column_cells[:COL_WIDTH_SAMPLE_ROWS]:

            if cell.value is None:
                continue

            max_length = max(max_length, len(str(cell.value)))

        width = max(MIN_COL_WIDTH, min(max_length + 2, MAX_COL_WIDTH))

        ws.column_dimensions[column_letter].width = width

    # --------------------------------------------------------
    # 保存（临时文件 + os.replace 原子替换）
    # --------------------------------------------------------

    tmp_file = output_file.with_name(output_file.name + ".tmp")
    wb.save(tmp_file)
    os.replace(tmp_file, output_file)

    print(MSG_DONE.format(output_file))
    print()


# ==================== 交互输入 ====================
def get_input_with_default(prompt_text: str, default_value: str) -> str:
    """获取带默认值的用户输入。"""
    user_input = input(f"{prompt_text} (默认: {default_value}): ").strip()
    return user_input if user_input else str(default_value)


def get_xml_file() -> Path:
    """确定 XML 文件：命令行参数优先，否则交互询问（带默认值）。"""
    if len(sys.argv) >= 2:
        xml_file = Path(sys.argv[1])
        if not xml_file.exists():
            print(MSG_FILE_NOT_FOUND.format(xml_file))
            sys.exit(1)
        return xml_file

    while True:
        user_input = get_input_with_default(MSG_PROMPT_XML, str(DEFAULT_MAME_XML))
        xml_file = Path(user_input)
        if xml_file.exists():
            return xml_file
        print(MSG_FILE_NOT_FOUND.format(xml_file))


def choose_output_mode() -> str:
    """选择输出方式，返回 "1" 或 "2"。"""
    print(MSG_MODE_MENU)

    while True:
        choice = input(MSG_PROMPT_MODE).strip()
        if choice in ("1", "2"):
            return choice
        print(MSG_INVALID_CHOICE)


# ==================== 主程序 ====================
def main() -> None:
    _init_quit_handler()

    # --------------------------------------------------------
    # 参数检查（仅命令行模式需要）
    # --------------------------------------------------------

    if len(sys.argv) >= 2 and not Path(sys.argv[1]).exists():
        print(MSG_FILE_NOT_FOUND.format(sys.argv[1]))
        sys.exit(1)

    # --------------------------------------------------------
    # 标题
    # --------------------------------------------------------

    print()
    print(MSG_SEPARATOR)
    print(MSG_TITLE)
    print(MSG_SEPARATOR)
    print()

    # --------------------------------------------------------
    # 输入：XML 文件 + 输出模式（解析前完成所有交互）
    # --------------------------------------------------------

    xml_file = get_xml_file()
    mode = choose_output_mode()

    # --------------------------------------------------------
    # 解析 XML
    # --------------------------------------------------------

    rows = parse_mame_xml(xml_file)

    if not rows:
        print(MSG_NO_DATA)
        return

    output_dir = xml_file.parent

    # ========================================================
    # 模式 1：一个 XLSX
    # ========================================================

    if mode == "1":

        if len(sys.argv) >= 3:

            output_file = Path(sys.argv[2])

            if not output_file.is_absolute():
                output_file = output_dir / output_file

        else:

            output_file = output_dir / f"{xml_file.stem}_database.xlsx"

        create_xlsx(rows, output_file, "MAME Database")

    # ========================================================
    # 模式 2：四个 XLSX
    # ========================================================

    elif mode == "2":

        datasets = {
            "Arcade_ROM": [
                row for row in rows
                if row["category"] == "arcade" and row["type"] == "rom"
            ],
            "Arcade_CHD": [
                row for row in rows
                if row["category"] == "arcade" and row["type"] == "chd"
            ],
            "Software_ROM": [
                row for row in rows
                if row["category"] == "software" and row["type"] == "rom"
            ],
            "Software_CHD": [
                row for row in rows
                if row["category"] == "software" and row["type"] == "chd"
            ],
        }

        for name, dataset in datasets.items():

            if not dataset:
                print(MSG_SKIP_EMPTY.format(name))
                continue

            output_file = output_dir / f"{xml_file.stem}_{name}.xlsx"
            create_xlsx(dataset, output_file, name)

    # --------------------------------------------------------
    # 最终结果
    # --------------------------------------------------------

    print()
    print(MSG_SEPARATOR)
    print(MSG_ALL_DONE)
    print(MSG_SEPARATOR)
    print()


# ==================== 程序入口 ====================
if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except (OSError, AttributeError):
        pass  # 管道/重定向时不崩溃

    try:
        main()
    except KeyboardInterrupt:
        print(MSG_INTERRUPTED)
    except Exception as e:
        print(MSG_ERROR.format(e))
    finally:
        try:
            input(MSG_EXIT)
        except EOFError:
            pass  # 管道输入结束时不崩溃
